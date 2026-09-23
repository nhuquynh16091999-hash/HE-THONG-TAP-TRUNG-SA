# -*- coding: utf-8 -*-
"""Sinh 2 file mẫu giống thật để thử hệ thống khi chưa có file tuần này.

Cố tình gài đủ mọi loại lệch mà hệ thống phải bắt được — xem README mục
"File mẫu" để biết đáp án.
"""
from openpyxl import Workbook
from datetime import datetime, date
import os

HERE = os.path.dirname(os.path.abspath(__file__))

# ── File 1: chi phí thanh toán từ TKQC Facebook ───────────────────────────
wb = Workbook()
ws = wb.active
ws.title = "Payment activity"
ws.append(["Transaction ID", "Transaction Date", "Ad Account ID", "Ad Account Name",
           "Payment Method", "Amount", "Currency", "Status"])

FB = [
    ("TXN-9001", datetime(2026, 9, 1),  "act_256210853826298", "Sỹ Anh Taiwan 4", "Visa *4281", 12_500_000, "VND", "Paid"),
    ("TXN-9002", datetime(2026, 9, 1),  "act_355509637365145", "Sỹ Anh Taiwan 5", "Visa *4281",  8_200_000, "VND", "Paid"),
    ("TXN-9003", datetime(2026, 9, 2),  "act_256210853826298", "Sỹ Anh Taiwan 4", "Visa *4281", 15_000_000, "VND", "Paid"),
    ("TXN-9004", datetime(2026, 9, 2),  "act_620354173588998", "Sỹ Anh Taiwan 6", "Mastercard *7733", 6_750_000, "VND", "Paid"),
    ("TXN-9005", datetime(2026, 9, 3),  "act_355509637365145", "Sỹ Anh Taiwan 5", "Visa *4281",  9_900_000, "VND", "Paid"),
    ("TXN-9006", datetime(2026, 9, 3),  "act_256210853826298", "Sỹ Anh Taiwan 4", "Visa *4281", 11_300_000, "VND", "Failed"),
    ("TXN-9007", datetime(2026, 9, 4),  "act_256210853826298", "Sỹ Anh Taiwan 4", "Visa *4281", 11_300_000, "VND", "Paid"),
    ("TXN-9008", datetime(2026, 9, 4),  "act_9999999999999",   "TK LẠ KHÔNG RÕ",  "Visa *4281",  4_500_000, "VND", "Paid"),
    ("TXN-9009", datetime(2026, 9, 5),  "act_620354173588998", "Sỹ Anh Taiwan 6", "Mastercard *7733", 7_400_000, "VND", "Paid"),
    ("TXN-9010", datetime(2026, 9, 5),  "act_355509637365145", "Sỹ Anh Taiwan 5", "Visa *4281", 13_800_000, "VND", "Paid"),
    ("TXN-9011", datetime(2026, 9, 6),  "act_256210853826298", "Sỹ Anh Taiwan 4", "Visa *4281", 10_100_000, "VND", "Paid"),
    ("TXN-9012", datetime(2026, 9, 7),  "act_355509637365145", "Sỹ Anh Taiwan 5", "Visa *4281",  5_600_000, "VND", "Paid"),
]
for r in FB:
    ws.append(list(r))
for c in ws["B"][1:]:
    c.number_format = "dd/mm/yyyy"
wb.save(os.path.join(HERE, "MAU_chi-phi-tkqc-fb_2026-09-01_07.xlsx"))

# ── File 2: sao kê thẻ ngân hàng ──────────────────────────────────────────
wb2 = Workbook()
ws2 = wb2.active
ws2.title = "Sao ke"
ws2.append(["Ngày giao dịch", "Số tham chiếu", "Nội dung giao dịch",
            "Số tiền ghi nợ", "Số tiền ghi có", "Số dư"])

BANK = [
    (datetime(2026, 9, 1), "FT26090100123", "FACEBK *4XY7Q2M8 VISA*4281",              12_500_000, 0, 487_500_000),
    (datetime(2026, 9, 1), "FT26090100455", "FACEBK *K92LM3P1 VISA*4281",               8_200_000, 0, 479_300_000),
    (datetime(2026, 9, 2), "FT26090200781", "FACEBK *TT81ZQ44 VISA*4281",              15_450_000, 0, 463_850_000),
    (datetime(2026, 9, 2), "FT26090200912", "FACEBK *M2L8XX09 MASTER*7733",             6_750_000, 0, 457_100_000),
    (datetime(2026, 9, 3), "FT26090301044", "TT LUONG NHAN VIEN THANG 8",             120_000_000, 0, 337_100_000),
    (datetime(2026, 9, 4), "FT26090400220", "FACEBK *QQ44MM10 VISA*4281",               9_900_000, 0, 327_200_000),
    (datetime(2026, 9, 4), "FT26090400221", "FACEBK *QQ44MM10 VISA*4281",               9_900_000, 0, 317_300_000),
    (datetime(2026, 9, 5), "FT26090500310", "FACEBK *AB77CD21 VISA*4281",              11_300_000, 0, 306_000_000),
    (datetime(2026, 9, 5), "FT26090500311", "FACEBK *ZZ10PP93 VISA*9911",               4_500_000, 0, 301_500_000),
    (datetime(2026, 9, 5), "FT26090500402", "GOOGLE ADS 8Y2K1 VISA*4281",               3_200_000, 0, 298_300_000),
    (datetime(2026, 9, 6), "FT26090600155", "FACEBK *PL02QW77 MASTER*7733",             7_400_000, 0, 290_900_000),
    (datetime(2026, 9, 6), "FT26090600156", "FACEBK *RT91NB05 VISA*4281",              13_800_000, 0, 277_100_000),
    (datetime(2026, 9, 6), "FT26090600157", "PHI GIAO DICH QUOC TE FACEBK",                62_500, 0, 277_037_500),
    (datetime(2026, 9, 7), "FT26090700198", "FACEBK *YU33HH88 VISA*4281",              10_100_000, 0, 266_937_500),
    (datetime(2026, 9, 8), "FT26090800011", "FACEBK *NN55JJ22 VISA*4281",              21_000_000, 0, 245_937_500),
]
for r in BANK:
    ws2.append(list(r))
for c in ws2["A"][1:]:
    c.number_format = "dd/mm/yyyy"
wb2.save(os.path.join(HERE, "MAU_sao-ke-the_2026-09-01_08.xlsx"))

print("Đã tạo 2 file mẫu trong", HERE)
