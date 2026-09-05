#!/usr/bin/env python3
"""Update 'Hệ thống đối soát tự động.docx' with new Financial Control features."""

import sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from docx import Document
from docx.shared import Pt

DOCX_PATH = r"C:\Users\Windows 11 Pro\Desktop\Hệ thống đối soát tự động.docx"

doc = Document(DOCX_PATH)
paragraphs = doc.paragraphs

# ═══ UPDATE 1: Update Dashboard section descriptions ═══
for para in paragraphs:
    text = para.text.strip()

    if text.startswith("Overview (Tổng quan)"):
        para.clear()
        r = para.add_run("Overview (Tổng quan): ")
        r.bold = True
        para.add_run(
            "Hiển thị 4 chỉ số KPI: Tổng tiền chờ thanh toán (Pending Payout EUR), "
            "Tỷ lệ thanh toán thành công (Settlement Rate %), "
            "Số ngày chờ tiền trung bình (Avg Days to Payout), Nợ cũ TCE. "
            "Biểu đồ Cash Cycle Timeline (Order → Delivery → Payout) và "
            "Aging Breakdown phân nhóm đơn chưa thanh toán theo độ tuổi "
            "(0-7 ngày, 8-14 ngày, 15-30 ngày, 30+ ngày). "
            "Cảnh báo tự động khi Pending > EUR 5,000 hoặc có đơn quá hạn 30+ ngày."
        )

    elif text.startswith("Reconciliation (Đối soát)"):
        para.clear()
        r = para.add_run("Reconciliation (Đối soát): ")
        r.bold = True
        para.add_run(
            "Bảng so khớp dòng tiền dạng waterfall: COD Total → COD Received → "
            "FFM Cost Deducted → Net Received → Pending → Collection Rate%. "
            "Phân tích theo từng thị trường (RO/BG/SK/HR). "
            "Danh sách chi tiết đơn hàng chờ thanh toán (lọc bỏ đơn COD = 0). "
            "Phần TCE Legacy Debt theo dõi nợ cũ từ đối tác đã ngừng hợp tác."
        )

    elif text.startswith("Cash Flow (Dòng tiền)"):
        para.clear()
        r = para.add_run("Cash Flow (Dòng tiền): ")
        r.bold = True
        para.add_run(
            "Toàn bộ số liệu hiển thị bằng VNĐ "
            "(tỷ giá: 1 USD = 26.325 VNĐ, 1 RON = 5.946 VNĐ, 1 EUR = 30.667 VNĐ). "
            "6 KPI cards: Revenue, Ads Spend, COGS (Nhập hàng), Fulfillment, Total Cost, Net Profit. "
            "Biểu đồ Stacked Bar theo tuần (Revenue vs Ads + COGS + FFM) kèm đường "
            "Cumulative Net Profit. Bảng Weekly Breakdown chi tiết từng tuần với Margin%. "
            "Dữ liệu từ mart_performance_master (pipeline tự động)."
        )

    elif text.startswith("Markets (Thị trường)"):
        para.clear()
        r = para.add_run("Markets (Thị trường): ")
        r.bold = True
        para.add_run(
            "Bảng so sánh hiệu quả tài chính giữa 4 thị trường EU: "
            "Romania (RON), Bulgaria (EUR), Slovakia (EUR), Croatia (EUR). "
            "Chỉ số: Tổng đơn, Đơn giao thành công, Đơn hoàn, COD, Chi phí FFM, "
            "Net, Pending Payout, Số ngày chờ trung bình, Tỷ lệ thanh toán %."
        )

    elif text.startswith("Verification (Xác minh)"):
        para.clear()
        r = para.add_run("Verification (Xác minh): ")
        r.bold = True
        para.add_run(
            'Module đối soát chính với 3 nút hành động: '
            '(1) Sync from Email — tự động lấy Protokol từ finance@levelupvn.com qua Gmail API; '
            '(2) Resync Payout — đồng bộ lại dữ liệu thanh toán từ euShipments API + re-verify; '
            '(3) Check POS Status — kiểm tra đơn đã cập nhật trạng thái "Đã thu tiền" trên POS chưa. '
            'Danh sách Payment Protocols hiển thị theo thứ tự mới nhất, phân biệt '
            'Gross Payout vs Net (deducted). Quy trình duyệt 2 người với Transaction ID ngân hàng. '
            'Dữ liệu xác nhận lưu trong BigQuery (không mất khi deploy).'
        )

# ═══ UPDATE 2: Add new sections at the end ═══

# Section 7: Automation
doc.add_paragraph("\u2500" * 80)
p7 = doc.add_paragraph()
r7 = p7.add_run("7. Tự động hóa (Cron Jobs)")
r7.bold = True
r7.font.size = Pt(13)

doc.add_paragraph(
    "Hệ thống chạy tự động trên VPS theo lịch hàng ngày:"
)

tbl = doc.add_table(rows=4, cols=3)
tbl.style = "Normal Table"
for j, text in enumerate(["Giờ (VN)", "Tác vụ", "Mô tả"]):
    tbl.rows[0].cells[j].text = text
    for run in tbl.rows[0].cells[j].paragraphs[0].runs:
        run.bold = True

tbl.rows[1].cells[0].text = "13:00"
tbl.rows[1].cells[1].text = "eu_shipment_sync.py --resync-all"
tbl.rows[1].cells[2].text = "Đồng bộ dữ liệu thanh toán (payout_date, payout_number) từ euShipments API"

tbl.rows[2].cells[0].text = "14:00"
tbl.rows[2].cells[1].text = "protokol_email_sync.py --reprocess"
tbl.rows[2].cells[2].text = "Quét email Protokol mới + re-verify + gửi Discord nếu có anomaly"

tbl.rows[3].cells[0].text = "Thứ 3 & Thứ 5"
tbl.rows[3].cells[1].text = "euShipments thanh toán"
tbl.rows[3].cells[2].text = "euShipments gửi email Protokol kèm XLSX/PDF. Hệ thống tự parse trong lần sync tiếp theo"

# Section 8: Discord
doc.add_paragraph("")
p8 = doc.add_paragraph()
r8 = p8.add_run("8. Thông báo Discord")
r8.bold = True
r8.font.size = Pt(13)

doc.add_paragraph(
    "Khi có Protokol mới hoặc phát hiện bất thường, hệ thống tự động gửi "
    "thông báo đến channel #đối-soát trên Discord:"
)
for item in [
    "Ngày + Số tiền EUR (hiển thị NET cho Compensation, GROSS cho Protokol thường)",
    "Link mở email Protokol gốc trong Gmail",
    "Link mở dashboard để duyệt",
    "Chi tiết anomaly (nếu có): AWB lạ, sai lệch số tiền, API chưa cập nhật",
    "Với Compensation: hiển thị breakdown COD Gross → Cấn trừ → Thực nhận",
]:
    doc.add_paragraph("  \u2022 " + item)

# Section 9: Data persistence
doc.add_paragraph("")
p9 = doc.add_paragraph()
r9 = p9.add_run("9. Lưu trữ dữ liệu")
r9.bold = True
r9.font.size = Pt(13)

doc.add_paragraph(
    "Dữ liệu xác nhận (Transaction ID, người duyệt, thời gian, dispute) "
    "được lưu trong BigQuery (bảng payout_protocols) — đảm bảo không mất "
    "dữ liệu khi deploy code mới hoặc restart server."
)
doc.add_paragraph(
    "Dữ liệu verification (items, anomalies, match results) lưu trong file "
    "JSON trên server — được tái tạo tự động bởi cron job hàng ngày."
)

# Section 10: Multi-team
doc.add_paragraph("")
p10 = doc.add_paragraph()
r10 = p10.add_run("10. Hỗ trợ đa team")
r10.bold = True
r10.font.size = Pt(13)

doc.add_paragraph(
    "Hệ thống phân biệt Protokol theo Client ID trong tên file "
    "(VD: Protokol-43472-3282.xlsx):"
)

tbl2 = doc.add_table(rows=5, cols=3)
tbl2.style = "Normal Table"
for j, text in enumerate(["Team", "Shop", "Client ID"]):
    tbl2.rows[0].cells[j].text = text
    for run in tbl2.rows[0].cells[j].paragraphs[0].runs:
        run.bold = True

data = [
    ("Stramark", "Aurelia Wear", "3282 & 3248"),
    ("Trendify", "Sorina", "3240 & 3283"),
    ("PiAlpha", "Xophia", "3239 & 3281"),
    ("T1", "Oddie House", "3284"),
]
for i, (team, shop, cid) in enumerate(data, 1):
    tbl2.rows[i].cells[0].text = team
    tbl2.rows[i].cells[1].text = shop
    tbl2.rows[i].cells[2].text = cid

doc.add_paragraph(
    "Hiện tại hệ thống chỉ xử lý Protokol của Stramark (Client 3282 & 3248). "
    "Khi vận hành ổn định sẽ mở rộng cho các team khác."
)

# Save
doc.save(DOCX_PATH)
print("Document updated successfully!")
